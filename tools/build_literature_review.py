#!/usr/bin/env python3
"""Build the Word review and its standalone recommendation tables.

Documentation-only dependency: python-docx==1.2.0. Experiment dependencies and
controls are not changed. Run from any directory with Python 3.10 or later.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/literature-review.md"
LEDGER = ROOT / "docs/literature_sources.tsv"
BLUE = "1F4E78"
GREY = "536579"
CITATION = re.compile(r"\[(\d+(?:\s*[,–-]\s*\d+)*)\]")
INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\)|\[\d+(?:\s*[,–-]\s*\d+)*\])")


def read_sources():
    with LEDGER.open(encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f, delimiter="\t"))
    assert [int(r["id"]) for r in rows] == list(range(1, len(rows) + 1))
    assert all(r["url"].startswith("https://") and r["title"] and r["access"] for r in rows)
    return {int(r["id"]): r for r in rows}


def citation_ids(text):
    ids = set()
    for match in CITATION.finditer(text):
        for part in re.split(r"\s*,\s*", match.group(1)):
            bounds = re.split("[–-]", part)
            start, end = int(bounds[0]), int(bounds[-1])
            assert start <= end
            ids.update(range(start, end + 1))
    return ids


def bookmark(paragraph, name, counter):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(counter))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(counter))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def link(paragraph, label, *, url=None, anchor=None, size=None):
    element = OxmlElement("w:hyperlink")
    if url:
        rid = paragraph.part.relate_to(
            url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
        )
        element.set(qn("r:id"), rid)
    if anchor:
        element.set(qn("w:anchor"), anchor)
    element.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    props.append(color)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(round(size * 2)))
        props.append(sz)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    text.set(qn("xml:space"), "preserve")
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def inline(paragraph, text, *, size=None, header=False):
    for token in INLINE.split(text):
        if not token:
            continue
        cite = CITATION.fullmatch(token)
        if cite and not header:
            first = re.match(r"\d+", cite.group(1)).group()
            link(paragraph, token, anchor=f"reference_{first}", size=size)
            continue
        explicit = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
        if explicit:
            label, target = explicit.groups()
            if target.startswith("https://"):
                link(paragraph, label, url=target, size=size)
            else:
                paragraph.add_run(label)
            continue
        run = paragraph.add_run(token[1:-1] if token.startswith("`") else token.strip("*") if token.startswith("**") else token)
        if token.startswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(size or 9.5)
        elif size:
            run.font.size = Pt(size)
        if token.startswith("**") or header:
            run.bold = True
        if header:
            run.font.color.rgb = RGBColor(255, 255, 255)


def page_size(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Cm(29.7 if landscape else 21)
    section.page_height = Cm(21 if landscape else 29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.7 if landscape else 2)
    section.header_distance = section.footer_distance = Cm(0.85)


def style_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for name, size in [("Title", 30), ("Subtitle", 15), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(13)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    for name in ["Header", "Footer"]:
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.size = Pt(8)
        doc.styles[name].font.color.rgb = RGBColor.from_string(GREY)
    section = doc.sections[0]
    page_size(section)
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.add_run("THE FEASIBLE RANK WINDOW  |  LITERATURE REVIEW")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Research synthesis • 17 October 2026  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    settings = OxmlElement("w:updateFields")
    settings.set(qn("w:val"), "true")
    doc.settings.element.append(settings)
    core = doc.core_properties
    core.title = "Literature Review: The Feasible Rank Window"
    core.subject = "LoRA rank, learning, retention, spectra and cost in four fixed decoder-only base models"
    core.author = "FYP research review"
    core.keywords = "LoRA, continual learning, BWT, intruder dimensions, hyperparameters, reproducibility"
    core.comments = "Generated from the review Markdown and audited source ledger. The review records a v2 baseline; selected recommendations are implemented in the separately documented v3 amendment."
    core.created = core.modified = datetime(2026, 10, 17, tzinfo=timezone.utc)


def set_cell(cell, text, *, header=False, shade=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    props = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side in ["top", "left", "bottom", "right"]:
        item = OxmlElement("w:" + side)
        item.set(qn("w:w"), "85" if side in {"top", "bottom"} else "95")
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    props.append(margins)
    if header or shade:
        fill = OxmlElement("w:shd")
        fill.set(qn("w:fill"), BLUE if header else "F0F4F8")
        props.append(fill)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.04
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = False
    inline(p, text, size=9, header=header)


def table(doc, rows, landscape):
    headers, *body = rows
    columns = len(headers)
    assert all(len(row) == columns for row in body), "Unequal table row widths"
    width = 25.7 if landscape else 17.0
    if columns == 5 and landscape:
        ratios = [0.08, 0.29, 0.22, 0.25, 0.16]
    elif columns == 5:
        ratios = [0.31, 0.15, 0.18, 0.16, 0.20]
    elif columns == 4:
        ratios = [0.17, 0.25, 0.29, 0.29]
    elif columns == 2:
        ratios = [0.35, 0.65]
    else:
        ratios = [1 / columns] * columns
    if headers[0] == "Fixed model":
        ratios = [0.18, 0.12, 0.29, 0.12, 0.29]
    t = doc.add_table(rows=1, cols=columns)
    t.style = "Table Grid"
    t.autofit = False
    for col, ratio in zip(t.columns, ratios):
        col.width = Cm(width * ratio)
    for row_index, values in enumerate([headers, *body]):
        row = t.rows[0] if row_index == 0 else t.add_row()
        tr_props = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_props.append(cant_split)
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_props.append(repeat)
        for cell, ratio, value in zip(row.cells, ratios, values):
            cell.width = Cm(width * ratio)
            set_cell(cell, value, header=row_index == 0, shade=row_index % 2 == 0)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.line_spacing = 0.3
    return t


def render_body(doc, markdown, heading_ids):
    lines = markdown.splitlines()
    i = 0
    landscape = False
    while i < len(lines):
        line = lines[i].strip()
        if not line or line.startswith("# "):
            i += 1
            continue
        if line in {"<!-- landscape -->", "<!-- portrait -->"}:
            target = "landscape" in line
            if target != landscape:
                sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
                page_size(sec, target)
                sec.different_first_page_header_footer = False
                landscape = target
            i += 1
            continue
        if line.startswith("<!--"):
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", x.replace(" ", "")) for x in cells):
                    rows.append(cells)
                i += 1
            table(doc, rows, landscape)
            continue
        if line.startswith("## ") or line.startswith("### "):
            depth = 1 if line.startswith("## ") else 2
            title = line.split(" ", 1)[1]
            p = doc.add_heading(title, level=depth)
            if title in heading_ids:
                bookmark(p, heading_ids[title], 1000 + list(heading_ids).index(title))
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.right_indent = Cm(0.3)
            r = p.add_run(line[2:])
            r.font.name = "Cambria Math"
            r.font.size = Pt(11)
            i += 1
            continue
        paragraph = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].lstrip().startswith(("#", "|", ">", "<!--")):
            paragraph.append(lines[i].strip())
            i += 1
        inline(doc.add_paragraph(), " ".join(paragraph))


def references(doc, sources):
    for key, source in sources.items():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        bookmark(p, f"reference_{key}", key)
        p.add_run(f"[{key}] {source['authors']} ({source['year']}). ")
        title = p.add_run(source["title"] + ". ")
        title.italic = True
        p.add_run(source["venue"] + ". ")
        link(p, "Primary source", url=source["url"], size=10)
        note = doc.add_paragraph("Access: " + source["access"])
        note.paragraph_format.left_indent = Cm(0.8)
        note.paragraph_format.space_after = Pt(5)
        for run in note.runs:
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(GREY)


def make_recommendations(markdown, sources):
    selected = markdown.split("## 9. Improvement tables:", 1)[1].split("## 10.", 1)[0]
    selected = selected.split("\n", 1)[1]
    selected = re.sub(r"<!-- (?:landscape|portrait) -->\n?", "", selected)
    out = "# Research improvements and candidate hyperparameters\n\n"
    out += "Extracted from the [full literature review](literature-review.md) and [Word review](../literature-review.docx), reviewed 17 October 2026. The four selected base models remain fixed. Selected measurement and LR/head-calibration proposals are implemented in [frw-v3](PROTOCOL_V3.md); the remaining entries are optional proposals. New/Extend statuses describe the reviewed v2 baseline.\n\n"
    out += selected.strip() + "\n\n## Sources cited in these tables\n\n"
    for key in sorted(citation_ids(selected)):
        s = sources[key]
        out += f"[{key}] {s['authors']} ({s['year']}). [{s['title']}]({s['url']}). {s['venue']}.\n\n"
    (ROOT / "docs/RESEARCH_IMPROVEMENTS.md").write_text(out, encoding="utf-8")


def build(output):
    sources = read_sources()
    markdown = SOURCE.read_text(encoding="utf-8")
    cited = citation_ids(markdown)
    assert cited <= set(sources), f"Unknown citation IDs: {cited - set(sources)}"
    assert set(sources) <= cited, f"Uncited bibliography entries: {set(sources) - cited}"
    recommendations = re.findall(r"^\| ([MH]\d{2}) / P[123] \|", markdown, re.M)
    assert len(recommendations) == len(set(recommendations)) == 28
    doc = Document()
    style_document(doc)
    cover = doc.add_paragraph()
    cover.paragraph_format.space_before = Cm(3.0)
    run = cover.add_run("LITERATURE REVIEW")
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(GREY)
    doc.add_paragraph("The Feasible\nRank Window", style="Title")
    doc.add_paragraph("Rank, learning and forgetting in compact\ndecoder-only language models", style="Subtitle")
    p = doc.add_paragraph("A critical review of the research evidence\nand 28 concrete design recommendations")
    p.paragraph_format.space_before = Cm(0.6)
    doc.add_paragraph("17 October 2026  •  62 source records  •  Review baseline: frw-v2; implementation: frw-v3")
    doc.add_paragraph("Fixed models: Llama-3.2-1B · Qwen2.5-1.5B · Granite-3.3-2B · Qwen2.5-3B")
    note = doc.add_paragraph("Recommendations are distinguished from published evidence. No new experimental results or optimal hyperparameters are claimed.")
    note.paragraph_format.space_before = Cm(0.8)
    note.runs[0].italic = True
    doc.add_page_break()
    doc.add_heading("Contents", level=1)
    headings = [line[3:] for line in markdown.splitlines() if line.startswith("## ")]
    heading_ids = {title: f"chapter_{i}" for i, title in enumerate(headings, 1)}
    for title in headings:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        link(p, title, anchor=heading_ids[title])
    doc.add_paragraph("Use the heading navigation pane to browse. Numbered citations link to the bibliography; each reference links to its primary source. Access qualifications are retained with the references.")
    doc.add_page_break()
    # Cover already contains the subtitle and provenance lines preceding the first section.
    body = "## " + markdown.split("## ", 1)[1]
    render_body(doc, body, heading_ids)
    references(doc, sources)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    make_recommendations(markdown, sources)
    print(f"Created {output}")
    print(f"Validated {len(sources)} sources, {len(cited)} cited IDs, {len(recommendations)} recommendations, {len(doc.tables)} tables, {len(doc.sections)} page sections.")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "literature-review.docx")
    args = parser.parse_args()
    build(args.output)


if __name__ == "__main__":
    main()
