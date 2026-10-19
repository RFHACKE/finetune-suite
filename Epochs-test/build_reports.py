"""Generate evidence-driven charts, a Word report and PowerPoint on CPU."""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import json
from pathlib import Path

import numpy as np

from rank_suite import config, utils
from rank_suite.protocol import digest
from epoch_study import MODELS, SEEDS, OBSERVED_EPOCHS

HERE = Path(__file__).resolve().parent
REFERENCES = [
    (
        "Hu et al. (2021/2022). LoRA: Low-Rank Adaptation of Large Language Models.",
        "https://arxiv.org/abs/2106.09685",
    ),
    (
        "Wang et al. (2018/2019). GLUE: A Multi-Task Benchmark and Analysis Platform for Natural Language Understanding.",
        "https://arxiv.org/abs/1804.07461",
    ),
    (
        "Dodge et al. (2020). Fine-Tuning Pretrained Language Models: Weight Initializations, Data Orders, and Early Stopping.",
        "https://arxiv.org/abs/2002.06305",
    ),
    ("Prechelt (1998). Early Stopping—But When?", "https://doi.org/10.1007/3-540-49430-8_3"),
    (
        "Hugging Face. Cosine schedule with warmup.",
        "https://huggingface.co/docs/transformers/main/en/main_classes/optimizer_schedules#transformers.get_cosine_schedule_with_warmup",
    ),
]


def mean_sd(values):
    a = np.asarray(values, dtype=float)
    return {
        "mean": float(a.mean()),
        "sd": float(a.std(ddof=1)) if len(a) > 1 else 0.0,
        "n": len(a),
        "values": a.tolist(),
    }


def read_rows(path):
    return [json.loads(line) for line in Path(path).read_text().splitlines() if line.strip()]


def prediction_metrics(path, expected, count):
    with np.load(path, allow_pickle=False) as f:
        labels, pred, prob, nll = (
            f[key] for key in ("labels", "predictions", "probabilities", "example_nll")
        )
    if (
        len(labels) != count
        or len(pred) != count
        or prob.shape != (count, 2)
        or len(nll) != count
        or not np.isfinite(prob).all()
        or not np.isfinite(nll).all()
        or not np.allclose(prob.sum(axis=1), 1, atol=1e-5)
        or not np.array_equal(pred, prob.argmax(axis=1))
    ):
        raise ValueError(f"Invalid prediction arrays: {path}")
    if (
        abs(float(nll.mean()) - expected["nll"]) > 1e-5
        or abs(float((labels == pred).mean()) - expected["accuracy"]) > 1e-12
        or not np.allclose(-np.log(np.maximum(prob[np.arange(count), labels], 1e-38)), nll, atol=1e-5)
    ):
        raise ValueError(f"Prediction/metric mismatch: {path}")
    wrong = labels != pred
    return {
        "brier": float(np.square(prob - np.eye(2)[labels]).sum(axis=1).mean()),
        "wrong_confidence": float(prob.max(axis=1)[wrong].mean()) if wrong.any() else None,
        "wrong_count": int(wrong.sum()),
    }


def analyze(root):
    root = Path(root)
    control = json.loads((root / "protocol.json").read_text())
    grid = json.loads((root / "expected_grid.json").read_text())
    sha = digest(control)
    records, missing, failures, rejected, run_times = [], [], [], [], []
    for identity in grid:
        run_root = root / "models" / identity["model"] / "runs" / identity["run_id"]
        path = run_root / "result.json"
        if not path.exists():
            missing.append(identity["run_id"])
            continue
        result = json.loads(path.read_text())
        if result.get("status") != "completed":
            failures.append(
                {"run_id": identity["run_id"], "status": result.get("status"), "error": result.get("error")}
            )
            continue
        try:
            if result.get("protocol_sha256") != sha or any(result.get(k) != v for k, v in identity.items()):
                raise ValueError("Identity/protocol mismatch")
            observations = read_rows(run_root / result["epoch_metrics"])
            if [r["epoch"] for r in observations] != list(OBSERVED_EPOCHS):
                raise ValueError("Expected exactly six ordered epoch observations")
            local = []
            step_rows = read_rows(run_root / result["attempt"] / "train_steps.jsonl")
            if len(step_rows) != 936 or [r["step"] for r in step_rows] != list(range(1, 937)):
                raise ValueError("Training did not record the planned 936 optimizer steps")
            for row in observations:
                if row["run_id"] != identity["run_id"] or row["task"] != "sst2":
                    raise ValueError("Observation identity/task mismatch")
                flat = {
                    "run_id": identity["run_id"],
                    "model": identity["model"],
                    "seed": identity["seed"],
                    "epoch": row["epoch"],
                    "rank": identity["rank"],
                    "lr": identity["lr"],
                    "head_lr_multiplier": identity["head_lr_multiplier"],
                    "online_train_loss": row["online_train_loss"],
                    "nll_gap": row["nll_gap"],
                    "accuracy_gap": row["accuracy_gap"],
                }
                for split, count in (("train", 2490), ("validation", 872)):
                    if row[split]["count"] != count:
                        raise ValueError("Unexpected evaluation count")
                    extra = prediction_metrics(run_root / row[split]["predictions"], row[split], count)
                    flat.update({f"{split}_{key}": row[split][key] for key in ("nll", "accuracy", "count")})
                    flat.update({f"{split}_{key}": value for key, value in extra.items()})
                steps = [s for s in step_rows if s["epoch"] == row["epoch"]]
                flat["mean_adapter_lr"] = float(np.mean([s["lr"] for s in steps]))
                flat["last_adapter_lr"] = steps[-1]["lr"]
                local.append(flat)
            records.extend(local)
            run_times.append(
                {
                    "run_id": identity["run_id"],
                    "model": identity["model"],
                    "seed": identity["seed"],
                    "elapsed_s": result["elapsed_s"],
                    "peak_vram_gb": result["peak_vram_gb"],
                }
            )
        except (ValueError, KeyError, OSError) as exc:
            rejected.append({"run_id": identity["run_id"], "reason": str(exc)})
    complete = len(run_times) == len(grid) and not rejected and not failures
    models = {}
    keys = (
        "train_nll",
        "validation_nll",
        "train_accuracy",
        "validation_accuracy",
        "nll_gap",
        "accuracy_gap",
        "validation_brier",
        "mean_adapter_lr",
    )
    for model in MODELS:
        selected = [r for r in records if r["model"] == model]
        epochs = {}
        for epoch in OBSERVED_EPOCHS:
            rows = [r for r in selected if r["epoch"] == epoch]
            if rows:
                epochs[str(epoch)] = {
                    "seeds": [r["seed"] for r in rows],
                    **{key: mean_sd([r[key] for r in rows]) for key in keys},
                }
        model_complete = len(selected) == len(SEEDS) * 6
        reference = {r["seed"]: r for r in selected if r["epoch"] == 3}
        comparisons = []
        for epoch in (4, 5, 6):
            paired = [
                {
                    "seed": r["seed"],
                    "train_nll_change": r["train_nll"] - reference[r["seed"]]["train_nll"],
                    "validation_nll_change": r["validation_nll"] - reference[r["seed"]]["validation_nll"],
                    "accuracy_change_pp": 100
                    * (r["validation_accuracy"] - reference[r["seed"]]["validation_accuracy"]),
                }
                for r in selected
                if r["epoch"] == epoch and r["seed"] in reference
            ]
            if paired:
                changes = {
                    key: mean_sd([r[key] for r in paired])
                    for key in ("train_nll_change", "validation_nll_change", "accuracy_change_pp")
                }
                count = sum(r["train_nll_change"] < 0 and r["validation_nll_change"] > 0 for r in paired)
                comparisons.append(
                    {
                        "epoch": epoch,
                        "reference": 3,
                        "n_pairs": len(paired),
                        **changes,
                        "paired_values": paired,
                        "overfitting_direction_seeds": count,
                        "mean_overfitting_pattern": changes["train_nll_change"]["mean"] < 0
                        and changes["validation_nll_change"]["mean"] > 0,
                        "consistent_descriptive_pattern": count == len(SEEDS)
                        and changes["validation_nll_change"]["mean"] > 0.02,
                    }
                )
        best_nll = (
            min(epochs, key=lambda e: (epochs[e]["validation_nll"]["mean"], int(e)))
            if model_complete
            else None
        )
        best_acc = (
            max(epochs, key=lambda e: (epochs[e]["validation_accuracy"]["mean"], -int(e)))
            if model_complete
            else None
        )
        models[model] = {
            "complete": model_complete,
            "epochs": epochs,
            "best_observed_nll_epoch": int(best_nll) if best_nll else None,
            "best_observed_accuracy_epoch": int(best_acc) if best_acc else None,
            "epoch3_lowest_mean_nll": int(best_nll) == 3 if best_nll else None,
            "comparisons_to_epoch3": comparisons,
        }
    supported = (
        all(
            m["epoch3_lowest_mean_nll"]
            and len(m["comparisons_to_epoch3"]) == 3
            and all(c["consistent_descriptive_pattern"] for c in m["comparisons_to_epoch3"])
            for m in models.values()
        )
        if complete
        else None
    )
    summary = {
        "schema": "epoch-trajectory-analysis-v1",
        "complete": complete,
        "expected_runs": len(grid),
        "completed_runs": len(run_times),
        "epoch_observations": len(records),
        "scope": control["scope"],
        "protocol_sha256": sha,
        "models": models,
        "missing": missing,
        "failures": failures,
        "rejected": rejected,
        "epoch3_hypothesis_supported_descriptively": supported,
        "independent_epoch_budget_comparison": False,
        "created_utc": datetime.now(timezone.utc).isoformat(),
    }
    return summary, records, run_times, control


def verdict(summary):
    if not summary["complete"]:
        return f"Study incomplete: {summary['completed_runs']}/{summary['expected_runs']} runs passed validation. No overall epoch conclusion is declared."
    count = sum(m["epoch3_lowest_mean_nll"] for m in summary["models"].values())
    if summary["epoch3_hypothesis_supported_descriptively"]:
        return "Epoch 3 has the lowest observed mean validation NLL in all four models; epochs 4–6 show consistent descriptive overfitting versus epoch 3. The hypothesis is supported within these six-epoch SST-2 trajectories."
    return (
        f"Epoch 3 has the lowest observed mean validation NLL in {count}/4 models. "
        "The full claim that epoch 3 is best for every model and all later epochs consistently overfit is not supported."
    )


def model_statement(model, data):
    if not data["complete"]:
        return f"{config.MODELS[model].label}: incomplete; no preferred epoch declared."
    later = [str(c["epoch"]) for c in data["comparisons_to_epoch3"] if c["consistent_descriptive_pattern"]]
    return (
        f"{config.MODELS[model].label}: lowest mean NLL at epoch {data['best_observed_nll_epoch']}; "
        f"highest mean accuracy at epoch {data['best_observed_accuracy_epoch']}. "
        f"Consistent overfitting versus epoch 3 at: {', '.join(later) or 'none'}."
    )


def fmt(stat, scale=1, digits=3):
    return f"{stat['mean'] * scale:.{digits}f} ± {stat['sd'] * scale:.{digits}f}"


def figures(root, summary, records):
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update(
        {
            "font.size": 11,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.grid": True,
            "grid.alpha": 0.2,
            "figure.facecolor": "white",
        }
    )
    dest = root / "reports" / "figures"
    dest.mkdir(parents=True, exist_ok=True)
    paths = {}

    def finish(fig, name):
        fig.tight_layout()
        for suffix in ("png", "pdf"):
            fig.savefig(dest / f"{name}.{suffix}", dpi=180, bbox_inches="tight")
        plt.close(fig)
        paths[name] = dest / f"{name}.png"

    def curve(ax, model, key, label, color, scale=1):
        groups = summary["models"][model]["epochs"]
        xs = [e for e in OBSERVED_EPOCHS if str(e) in groups]
        ys = [groups[str(e)][key]["mean"] * scale for e in xs]
        sd = [groups[str(e)][key]["sd"] * scale for e in xs]
        ax.errorbar(xs, ys, yerr=sd, marker="o", capsize=4, label=label, color=color, linewidth=2)
        for si, seed in enumerate(SEEDS):
            dots = [r for r in records if r["model"] == model and r["seed"] == seed]
            ax.scatter(
                [r["epoch"] + (si - 1) * 0.045 for r in dots],
                [r[key] * scale for r in dots],
                color=color,
                s=15,
                alpha=0.5,
            )
        ax.set_xticks(OBSERVED_EPOCHS)
        ax.axvline(3, color="gray", ls=":", alpha=0.6)
        ax.set_xlabel("Epoch within continuous six-epoch training")

    fig, axes = plt.subplots(2, 2, figsize=(12, 7))
    for ax, model in zip(axes.flat, MODELS):
        curve(ax, model, "validation_nll", "Full validation NLL", "#ce6723")
        ax.set_title(config.MODELS[model].label)
        ax.set_ylabel("Validation NLL (lower is better)")
    fig.suptitle("Validation NLL by epoch: mean ± sample SD; dots = individual seeds", fontsize=14)
    finish(fig, "validation-overview")
    fig, axes = plt.subplots(2, 2, figsize=(12, 7))
    for ax, model in zip(axes.flat, MODELS):
        curve(ax, model, "validation_accuracy", "Full validation accuracy", "#16778a", 100)
        ax.set_title(config.MODELS[model].label)
        ax.set_ylabel("Validation accuracy (%)")
    finish(fig, "accuracy-overview")
    for model in MODELS:
        fig, axes = plt.subplots(2, 2, figsize=(12, 7))
        for split, label, color in (
            ("train", "Full train", "#16778a"),
            ("validation", "Full validation", "#ce6723"),
        ):
            curve(axes[0, 0], model, f"{split}_nll", label, color)
            curve(axes[0, 1], model, f"{split}_accuracy", label, color, 100)
        axes[0, 0].set_ylabel("NLL")
        axes[0, 1].set_ylabel("Accuracy (%)")
        axes[0, 0].legend()
        axes[0, 1].legend()
        curve(axes[1, 0], model, "nll_gap", "Validation NLL − train NLL", "#6855a5")
        axes[1, 0].set_ylabel("NLL generalization gap")
        axes[1, 0].legend(fontsize=9)
        pairs = summary["models"][model]["comparisons_to_epoch3"]
        for key, label, color in (
            ("train_nll_change", "Train change", "#16778a"),
            ("validation_nll_change", "Validation change", "#ce6723"),
        ):
            axes[1, 1].errorbar(
                [p["epoch"] for p in pairs],
                [p[key]["mean"] for p in pairs],
                yerr=[p[key]["sd"] for p in pairs],
                marker="o",
                capsize=4,
                label=label,
                color=color,
            )
        axes[1, 1].axhline(0, color="gray", lw=1)
        axes[1, 1].set_xticks((4, 5, 6))
        axes[1, 1].set_xlabel("Later epoch compared with epoch 3")
        axes[1, 1].set_ylabel("Paired NLL change (mean ± SD)")
        axes[1, 1].legend(fontsize=9)
        fig.suptitle(f"{config.MODELS[model].label}: full-set epoch measurements, rank 16", fontsize=14)
        finish(fig, f"{model}-epochs")
    matrix = np.full((4, 6), np.nan)
    for i, model in enumerate(MODELS):
        groups = summary["models"][model]["epochs"]
        if "3" in groups:
            for j, epoch in enumerate(OBSERVED_EPOCHS):
                if str(epoch) in groups:
                    matrix[i, j] = (
                        groups[str(epoch)]["validation_nll"]["mean"] - groups["3"]["validation_nll"]["mean"]
                    )
    finite = np.abs(matrix[np.isfinite(matrix)])
    limit = max(float(finite.max()) if len(finite) else 0.1, 0.01)
    fig, ax = plt.subplots(figsize=(11, 4))
    heat = ax.imshow(np.ma.masked_invalid(matrix), cmap="RdBu_r", vmin=-limit, vmax=limit, aspect="auto")
    ax.set_yticks(range(4), [config.MODELS[m].label for m in MODELS])
    ax.set_xticks(range(6), OBSERVED_EPOCHS)
    ax.set_xlabel("Observed epoch in six-epoch fits")
    for (i, j), v in np.ndenumerate(matrix):
        if np.isfinite(v):
            ax.text(
                j,
                i,
                f"{v:+.3f}",
                ha="center",
                va="center",
                color="white" if abs(v) > 0.6 * limit else "black",
            )
    ax.grid(False)
    ax.set_title("Validation NLL change versus epoch 3 (negative = better; positive = worse)")
    fig.colorbar(heat, ax=ax, label="NLL change")
    finish(fig, "comparison-to-epoch3")
    fig, axes = plt.subplots(1, 2, figsize=(12, 4))
    colors = ("#16778a", "#ce6723", "#6855a5", "#70832b")
    for model, color in zip(MODELS, colors):
        curve(axes[0], model, "validation_brier", config.MODELS[model].label, color)
        curve(axes[1], model, "mean_adapter_lr", config.MODELS[model].label, color)
    axes[0].set_ylabel("Validation Brier score (lower is better)")
    axes[1].set_ylabel("Mean adapter learning rate within epoch")
    for ax in axes:
        ax.legend(fontsize=8)
    finish(fig, "calibration-and-learning-rate")
    return paths


def word_report(root, summary, control, paths):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("15384D")
    doc.add_heading("Epochs, generalization and overfitting", 0)
    doc.add_paragraph("SST-2 · LoRA rank 16 · four models · three seeds", "Subtitle")
    doc.add_paragraph(
        f"Generated {summary['created_utc']}\n{summary['completed_runs']}/{summary['expected_runs']} complete runs; {summary['epoch_observations']} epoch observations\nProtocol SHA-256: {summary['protocol_sha256']}"
    )
    doc.add_heading("Findings", 1)
    doc.add_paragraph(verdict(summary))
    for model, result in summary["models"].items():
        doc.add_paragraph(model_statement(model, result))
    doc.add_heading("Question and design", 1)
    doc.add_paragraph(
        "We test whether validation performance is best at epoch 3 and worsens at epochs 4–6. This is a hypothesis, not an assumed result. Each of four pinned pretrained models is fine-tuned once per seed (42, 123, 456), continuously for six epochs: 12 runs and 72 epoch observations. After every epoch, all 2,490 fixed training examples and all 872 official SST-2 validation examples are evaluated. No early stopping, adaptive tuning, checkpoint selection, or discarded seeds are used."
    )
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    for c, t in zip(table.rows[0].cells, ("Model", "Adapter LR", "Head LR multiplier")):
        c.text = t
    for model in MODELS:
        for c, t in zip(
            table.add_row().cells,
            (
                config.MODELS[model].label,
                f"{control['learning_rates'][model]:.3g}",
                f"{control['head_lr_multipliers'][model]:g}",
            ),
        ):
            c.text = t
    doc.add_paragraph(
        "Fixed controls: q/v LoRA rank 16, alpha 32; AdamW, FP32 trainable parameters and moments, BF16 compute; effective/micro batch 16; evaluation batch 32; maximum length 256; weight decay 0; dropout 0; clipping 1; 3% warmup then cosine decay. These LR/head policies were best among the project's earlier tested candidates [1]. The dataset and all model revisions are pinned [2]."
    )
    doc.add_heading("Interpretation and schedule limitation", 2)
    doc.add_paragraph(
        "Negative log-likelihood (NLL) measures the mean negative log probability assigned to the correct label; lower is better. Overfitting is suggested when training NLL falls while validation NLL rises [4]. A larger gap alone is insufficient. Accuracy can remain stable even as wrong predictions become more confident; Brier score and per-example probabilities provide additional context. Brier score here is the mean sum of squared probability errors across both classes (range 0–2)."
    )
    doc.add_paragraph(
        "The same optimizer and six-epoch cosine schedule continue across observations. Epoch 3 here therefore differs from epoch 3 of the original three-epoch schedule, which decays faster [5]. This study describes learning trajectories and possible stopping points under the six-epoch schedule. It does not establish the best independent epoch budget or prove the original three-epoch policy optimal."
    )
    doc.add_paragraph(
        "All observations restore RNG state and never change training. Means and sample standard deviations are computed over the three seeds; error bars are SD, not confidence intervals. Individual seeds are also shown. Comparisons at epochs 4–6 are paired with epoch 3 within each seed. The predeclared consistency label requires falling train NLL and rising validation NLL in all three paired seeds plus mean validation NLL rise >0.02. This is a descriptive rule, not a significance test. Seed variability motivates retaining every replication [3]."
    )
    for name, caption in (
        ("validation-overview", "Full validation NLL: mean ± sample SD, with individual seeds."),
        ("accuracy-overview", "Full validation accuracy across epochs."),
        ("comparison-to-epoch3", "NLL changes relative to epoch 3. Blue = better; red = worse."),
    ):
        doc.add_picture(str(paths[name]), width=Inches(6.7))
        doc.add_paragraph(caption, "Caption")
    for model, result in summary["models"].items():
        doc.add_page_break()
        doc.add_heading(config.MODELS[model].label, 1)
        doc.add_paragraph(model_statement(model, result))
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Shading Accent 1"
        for c, t in zip(
            table.rows[0].cells, ("Epoch", "Train NLL", "Validation NLL", "Validation acc. (%)", "NLL gap")
        ):
            c.text = t
        for epoch, group in result["epochs"].items():
            for c, t in zip(
                table.add_row().cells,
                (
                    epoch,
                    fmt(group["train_nll"]),
                    fmt(group["validation_nll"]),
                    fmt(group["validation_accuracy"], 100, 2),
                    fmt(group["nll_gap"]),
                ),
            ):
                c.text = t
        doc.add_paragraph(
            "Mean ± sample SD across seeds. Each row is observed within continuing six-epoch fits.", "Caption"
        )
        doc.add_picture(str(paths[f"{model}-epochs"]), width=Inches(6.7))
        doc.add_heading("Changes after epoch 3", 2)
        for pair in result["comparisons_to_epoch3"]:
            doc.add_paragraph(
                f"Epoch {pair['epoch']} versus 3: train NLL {fmt(pair['train_nll_change'])}; validation NLL {fmt(pair['validation_nll_change'])}; accuracy {fmt(pair['accuracy_change_pp'], digits=2)} percentage points. Overfitting direction: {pair['overfitting_direction_seeds']}/{pair['n_pairs']} seeds. Consistent descriptive pattern: {'yes' if pair['consistent_descriptive_pattern'] else 'no'}."
            )
    doc.add_heading("Confidence and the learning-rate path", 1)
    doc.add_picture(str(paths["calibration-and-learning-rate"]), width=Inches(6.7))
    doc.add_paragraph(
        "Brier score measures probability error. The LR chart records actual mean step learning rates, showing the continuing six-epoch schedule. The seed-level CSV also retains wrong-prediction counts and mean confidence on wrong predictions."
    )
    doc.add_heading("Limitations and conclusion", 1)
    doc.add_paragraph(verdict(summary))
    doc.add_paragraph(
        "This is an exploratory SST-2 study using a small fixed training subset, three seeds, one rank, and previously selected learning rates. Previously inspected official validation data is reused, so it is not an untouched test-set confirmation. No universal optimum, statistical significance, or optimum for the full RTE → MNLI → SST-2 sequence is claimed. The observations within each run are correlated; 72 observations are not 72 independent replications. Model and metric preferences may differ. Original experiments and global defaults remain separate and unchanged."
    )
    if not summary["complete"]:
        doc.add_paragraph(
            f"Missing runs: {len(summary['missing'])}; failed/incomplete: {len(summary['failures'])}; rejected: {len(summary['rejected'])}. Exact identifiers appear in summary.json. Missing values are not imputed."
        )
    doc.add_heading("Reproducibility and evidence", 1)
    doc.add_paragraph(
        "JSON/CSV tables, all per-epoch full-set predictions/probabilities/NLL, optimizer-step losses and learning rates, data indices, initialization hashes, pinned source, hardware/software records, and failure logs are retained. Model weights, adapter weights, projection tensors and optimizer states are not exported. Word and PowerPoint derive from the same summary; PNG/PDF charts and CPU-only regeneration code are supplied."
    )
    doc.add_heading("References", 1)
    for i, (title, url) in enumerate(REFERENCES, 1):
        doc.add_paragraph(f"[{i}] {title}\n{url}")
    section.footer.paragraphs[
        0
    ].text = f"Separate SST-2 epoch study | protocol {summary['protocol_sha256'][:12]}"
    doc.save(root / "reports" / "epoch-study.docx")


def presentation(root, summary, paths):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)

    def text(slide, value, x, y, w, h, size=23, bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        for i, line in enumerate(value.split("\n")):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = line
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = RGBColor(21, 56, 77)
            p.space_after = Pt(12)

    def slide(title, body=None, chart=None, note=""):
        item = deck.slides.add_slide(deck.slide_layouts[6])
        text(item, title, 0.55, 0.25, 12.2, 0.7, 29, True)
        if body:
            text(item, body, 0.75, 1.3, 11.8, 5.4)
        if chart:
            from PIL import Image

            with Image.open(chart) as pic:
                w, h = pic.size
            scale = min(12.1 / w, 5.55 / h)
            w, h = w * scale, h * scale
            item.shapes.add_picture(
                str(chart), Inches((13.333 - w) / 2), Inches(1.0), width=Inches(w), height=Inches(h)
            )
        text(item, note, 0.65, 6.7, 12, 0.48, 12)
        text(item, f"Epochs-test | SST-2 | {len(deck.slides)}", 0.65, 7.15, 12, 0.2, 9)

    slide(
        "Epochs, generalization and overfitting",
        f"LoRA rank 16 • four fixed models • three seeds\n12 continuous six-epoch runs, 72 epoch measurements\n\n{verdict(summary)}",
        note="A test of the epoch-3 hypothesis; conclusions follow the observations.",
    )
    slide(
        "Design: evaluate while training continues",
        "• Train epoch 1 → evaluate → epoch 2 → evaluate → … → epoch 6.\n• Full 2,490-example training and 872-example validation evaluation each epoch.\n• Four models × seeds 42, 123, 456 = 12 independent runs.\n• Rank 16; selected model-specific learning rates; default training controls.\n• Same optimizer and six-epoch cosine schedule throughout; no early stopping.",
        note="Epoch measurements within a run are correlated, not independent replications.",
    )
    slide(
        "What would support overfitting?",
        "• Training NLL decreases: closer fit to training examples.\n• Validation NLL increases: worse correct-label probabilities on held-out data.\n• Accuracy can stay stable while wrong predictions become more confident.\n• Our descriptive consistency rule requires that pattern in all three seeds and mean validation NLL rise >0.02.\n• Compare epochs 4–6 with epoch 3 within each seed.",
        note="A widening gap alone is insufficient. The rule is not a statistical significance test.",
    )
    slide(
        "Validation NLL across epochs",
        chart=paths["validation-overview"],
        note="Lower is better. Error bars = sample SD; dots = seeds; dotted line = epoch-3 reference.",
    )
    slide(
        "Validation accuracy across epochs",
        chart=paths["accuracy-overview"],
        note="Accuracy and NLL can prefer different epochs; both are reported.",
    )
    slide(
        "How much does each epoch differ from epoch 3?",
        chart=paths["comparison-to-epoch3"],
        note="Blue/negative = lower NLL; red/positive = higher NLL. Check training loss before calling a change overfitting.",
    )
    for model, result in summary["models"].items():
        slide(
            f"{config.MODELS[model].label}: epoch comparisons",
            chart=paths[f"{model}-epochs"],
            note=model_statement(model, result),
        )
    slide(
        "Probability quality and learning rates",
        chart=paths["calibration-and-learning-rate"],
        note="Brier score: lower is better. The observed LR path belongs to the six-epoch schedule.",
    )
    slide(
        "Findings by model",
        "\n".join(model_statement(m, r) for m, r in summary["models"].items()),
        note="Best means lowest observed mean NLL or highest observed accuracy among these six epochs.",
    )
    slide(
        "Conclusion and limits",
        verdict(summary)
        + "\n\n• SST-2 only, fixed small training subset, three seeds, one rank.\n• Previously inspected validation data; no untouched test-set claim.\n• A six-epoch trajectory does not prove the original three-epoch schedule optimal.\n• Original rank-sweep data and defaults remain unchanged.",
    )
    slide(
        "Evidence retained for later analysis",
        "• Word report, this presentation, PNG/PDF charts, JSON and CSV.\n• Full-set per-epoch probabilities, labels, NLL, accuracy and gaps.\n• Step-level losses and actual learning rates; seed-level comparisons.\n• Pinned source/data revisions, selected indices and failure records.\n• No large model weights; CPU-only reanalysis is supported.",
    )
    slide(
        "References", "\n".join(f"[{i}] {title}\n{url}" for i, (title, url) in enumerate(REFERENCES[:3], 1))
    )
    slide(
        "References (continued)",
        "\n".join(f"[{i}] {title}\n{url}" for i, (title, url) in enumerate(REFERENCES[3:], 4)),
    )
    deck.save(root / "reports" / "epoch-study.pptx")


def write_csv(path, rows):
    if rows:
        with Path(path).open("w", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)


def build(root):
    root = Path(root)
    (root / "reports").mkdir(parents=True, exist_ok=True)
    summary, records, run_times, control = analyze(root)
    utils.save_json(root / "reports" / "summary.json", summary)
    write_csv(root / "reports" / "epoch-metrics.csv", records)
    write_csv(root / "reports" / "run-times.csv", run_times)
    comparisons = [
        {
            "model": m,
            "epoch": p["epoch"],
            "reference": 3,
            "n_pairs": p["n_pairs"],
            "train_nll_change": p["train_nll_change"]["mean"],
            "validation_nll_change": p["validation_nll_change"]["mean"],
            "accuracy_change_pp": p["accuracy_change_pp"]["mean"],
            "overfitting_direction_seeds": p["overfitting_direction_seeds"],
            "consistent_descriptive_pattern": p["consistent_descriptive_pattern"],
        }
        for m, data in summary["models"].items()
        for p in data["comparisons_to_epoch3"]
    ]
    write_csv(root / "reports" / "comparisons-to-epoch3.csv", comparisons)
    paths = figures(root, summary, records)
    word_report(root, summary, control, paths)
    presentation(root, summary, paths)
    report = {
        "complete": summary["complete"],
        "completed_runs": summary["completed_runs"],
        "expected_runs": summary["expected_runs"],
        "epoch_observations": summary["epoch_observations"],
        "verdict": verdict(summary),
        "figures": [str(p.relative_to(root)) for p in paths.values()],
        "documents": ["reports/epoch-study.docx", "reports/epoch-study.pptx"],
    }
    utils.save_json(root / "reports" / "manifest.json", report)
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=HERE / "results")
    print(json.dumps(build(parser.parse_args().root), indent=2))
