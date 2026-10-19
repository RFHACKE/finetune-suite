"""Scientific preflight checks; synthetic fixtures stay in pytest temp folders."""

import json
from pathlib import Path
import sys
import zipfile

import numpy as np
import pytest
import torch

sys.path.insert(0, str(Path(__file__).resolve().parent))
from epoch_study import EPOCH_BUDGETS, SEEDS, expected_grid, observation_rng, protocol  # noqa: E402
from build_reports import analyze, scheduled_lr, seed_diagnostics, verdict  # noqa: E402
from rank_suite import utils  # noqa: E402
from rank_suite.protocol import digest  # noqa: E402


def test_exact_independent_grid_and_controls():
    grid = expected_grid("sst2")
    assert len(grid) == len({r["run_id"] for r in grid}) == 48
    assert SEEDS == (42, 123)
    assert all(r["training_epochs"] == r["epoch_budget"] and r["rank"] == 16 for r in grid)
    assert sum(r["training_epochs"] for r in grid) == 168
    assert protocol("sst2")["expected_optimizer_steps"] == 26208
    assert [int(np.ceil(0.03 * 156 * b)) for b in EPOCH_BUDGETS] == [5, 10, 15, 19, 24, 29]
    with pytest.raises(ValueError):
        protocol("full")


def test_observations_do_not_change_training(tmp_path, monkeypatch):
    from rank_suite import train

    monkeypatch.setattr(train, "forward_logits", lambda model, head, batch: head(model(batch["x"])))
    outputs, metrics = [], []
    for observe in (False, True):
        utils.seed_all(42)
        model, head = torch.nn.Linear(4, 3), torch.nn.Linear(3, 2)
        data = [
            {"x": torch.arange(4, dtype=torch.float32) / (i + 1), "labels": torch.tensor(i % 2)}
            for i in range(16)
        ]
        loader = torch.utils.data.DataLoader(
            data, batch_size=4, shuffle=True, generator=torch.Generator().manual_seed(19)
        )

        def callback(epoch, step, loss):
            with observation_rng("cpu"):
                metrics.append(
                    train.evaluate(
                        model,
                        head,
                        torch.utils.data.DataLoader(data, batch_size=4),
                        "cpu",
                        return_details=True,
                    )
                )
                torch.rand(7)

        train._train_task(
            model,
            head,
            list(model.parameters()) + list(head.parameters()),
            loader,
            6,
            1e-3,
            "adamw",
            "cpu",
            tmp_path / f"steps-{observe}.jsonl",
            0,
            "sst2",
            epoch_callback=callback if observe else None,
        )
        outputs.append([p.detach().clone() for p in list(model.parameters()) + list(head.parameters())])
    assert len(metrics) == 6
    assert all(torch.equal(a, b) for a, b in zip(*outputs))


def test_lr_checker_matches_actual_training_schedules(tmp_path, monkeypatch):
    from rank_suite import train

    monkeypatch.setattr(train, "forward_logits", lambda model, head, batch: head(model(batch["x"])))
    torch.set_num_threads(1)
    initial = []
    for budget in (1, 3, 6):
        utils.seed_all(42)
        model, head = torch.nn.Linear(2, 2), torch.nn.Linear(2, 2)
        initial.append(model.weight.detach().clone())
        data = [{"x": torch.tensor([0.5, 0.2]), "labels": torch.tensor(i % 2)} for i in range(2490)]
        loader = torch.utils.data.DataLoader(
            data, batch_size=16, shuffle=True, generator=torch.Generator().manual_seed(4202)
        )
        path = tmp_path / f"budget-{budget}.jsonl"
        train._train_task(
            model,
            head,
            list(model.parameters()) + list(head.parameters()),
            loader,
            budget,
            4e-4,
            "adamw",
            "cpu",
            path,
            0,
            "sst2",
            head_lr_multiplier=5.0,
        )
        rows = [json.loads(line) for line in path.read_text().splitlines()]
        assert len(rows) == 156 * budget
        assert all(
            np.isclose(r["lr"], scheduled_lr(i, budget, 4e-4), rtol=1e-12, atol=1e-16)
            for i, r in enumerate(rows)
        )
    assert all(torch.equal(initial[0], p) for p in initial[1:])


def fixture_cohort(root, best_budget):
    control, grid = protocol("sst2"), expected_grid("sst2")
    utils.save_json(root / "protocol.json", control)
    utils.save_json(root / "expected_grid.json", grid)
    for identity in grid:
        run = root / "models" / identity["model"] / "runs" / identity["run_id"]
        attempt = run / "attempts" / "fixture"
        attempt.mkdir(parents=True)
        budget, observations = identity["epoch_budget"], []
        for epoch in range(1, budget + 1):
            row = {**identity, "epoch": epoch, "task": "sst2", "online_train_loss": 0.4, "step": 156 * epoch}
            for split, count in (("train", 2490), ("validation", 872)):
                labels = np.arange(count, dtype=np.int64) % 2
                probability = (
                    0.7 + 0.04 * epoch
                    if split == "train"
                    else 0.7 + 0.04 * min(budget, best_budget) - 0.08 * max(0, budget - best_budget)
                )
                # Extremely good intermediate measurements must NOT affect endpoint ranking.
                if split == "validation" and epoch < budget:
                    probability = 0.97
                true_prob = np.full(count, probability, dtype=np.float32)
                true_prob[::10] = 0.25
                prob = np.empty((count, 2), dtype=np.float32)
                prob[np.arange(count), labels], prob[np.arange(count), 1 - labels] = true_prob, 1 - true_prob
                nll, pred = -np.log(true_prob), prob.argmax(axis=1)
                path = attempt / f"epoch{epoch}_{split}.npz"
                np.savez_compressed(
                    path, labels=labels, predictions=pred, probabilities=prob, example_nll=nll
                )
                row[split] = {
                    "count": count,
                    "nll": float(nll.mean()),
                    "accuracy": float((labels == pred).mean()),
                    "predictions": str(path.relative_to(run)),
                }
            row["nll_gap"] = row["validation"]["nll"] - row["train"]["nll"]
            row["accuracy_gap"] = row["train"]["accuracy"] - row["validation"]["accuracy"]
            observations.append(row)
        (attempt / "epoch_metrics.jsonl").write_text("".join(json.dumps(r) + "\n" for r in observations))
        steps = [
            {
                "step": i + 1,
                "epoch": i // 156 + 1,
                "task": "sst2",
                "loss": 0.4,
                "grad_norm": 0.5,
                "examples": 10 if i % 156 == 155 else 16,
                "lr": scheduled_lr(i, budget, identity["lr"]),
                "head_lr": scheduled_lr(i, budget, identity["lr"]) * identity["head_lr_multiplier"],
            }
            for i in range(156 * budget)
        ]
        (attempt / "train_steps.jsonl").write_text("".join(json.dumps(r) + "\n" for r in steps))
        utils.save_json(
            run / "result.json",
            {
                **identity,
                "status": "completed",
                "protocol_sha256": digest(control),
                "attempt": "attempts/fixture",
                "epoch_metrics": "attempts/fixture/epoch_metrics.jsonl",
                "elapsed_s": 1.0,
                "peak_vram_gb": 1.0,
                "initialization": {
                    "adapter_sha256": f"initial-{identity['model']}-{identity['seed']}",
                    "heads_sha256": f"heads-{identity['seed']}",
                },
                "final_adapter_sha256": f"updated-{identity['run_id']}",
            },
        )
    return grid


@pytest.mark.parametrize("best_budget", [3, 4])
def test_endpoint_only_conclusions_follow_data(tmp_path, best_budget):
    fixture_cohort(tmp_path, best_budget)
    summary, rows, times, _ = analyze(tmp_path)
    assert summary["complete"]
    assert len(rows) == 168 and len(times) == 48
    assert summary["endpoint_observations"] == 48
    assert summary["optimizer_steps"] == 26208
    assert all(m["best_observed_nll_budget"] == best_budget for m in summary["models"].values())
    assert all(
        g["validation_nll"]["n"] == 2 for m in summary["models"].values() for g in m["budgets"].values()
    )
    assert summary["epoch3_hypothesis_supported_descriptively"] is (best_budget == 3)
    if best_budget == 4:
        assert "not supported" in verdict(summary)


def test_missing_epoch_blocks_conclusion(tmp_path):
    first = fixture_cohort(tmp_path, 3)[0]
    path = (
        tmp_path
        / "models"
        / first["model"]
        / "runs"
        / first["run_id"]
        / "attempts/fixture/epoch_metrics.jsonl"
    )
    path.write_text("")
    summary, *_ = analyze(tmp_path)
    assert not summary["complete"]
    assert summary["epoch3_hypothesis_supported_descriptively"] is None
    assert summary["rejected"][0]["run_id"] == first["run_id"]


def test_wrong_schedule_is_rejected(tmp_path):
    first = fixture_cohort(tmp_path, 3)[0]
    path = (
        tmp_path / "models" / first["model"] / "runs" / first["run_id"] / "attempts/fixture/train_steps.jsonl"
    )
    rows = [json.loads(line) for line in path.read_text().splitlines()]
    rows[50]["lr"] = scheduled_lr(50, 6, first["lr"])
    path.write_text("".join(json.dumps(r) + "\n" for r in rows))
    summary, *_ = analyze(tmp_path)
    assert not summary["complete"]
    assert "schedule" in summary["rejected"][0]["reason"]


def test_joint_regression_is_not_mislabeled_as_simple_overfitting():
    base = {"model": "qwen2.5-3b", "seed": 42, "is_endpoint": True}
    diagnostic = seed_diagnostics(
        [
            {
                **base,
                "epoch_budget": 3,
                "train_nll": 0.03,
                "validation_nll": 0.3,
                "validation_accuracy": 0.95,
            },
            {
                **base,
                "epoch_budget": 6,
                "train_nll": 0.17,
                "validation_nll": 0.54,
                "validation_accuracy": 0.8,
            },
        ]
    )
    regression = diagnostic["accuracy_regressions"][0]
    assert regression["both_losses_increased"]
    assert not regression["loss_overfitting_direction"]
    assert regression["accuracy_change_pp"] == pytest.approx(-15)


def test_documents_open_and_preserve_contrary_findings(tmp_path):
    from build_reports import build
    from docx import Document
    from pptx import Presentation

    fixture_cohort(tmp_path, 4)
    result = build(tmp_path)
    assert result["complete"] and result["epoch_observations"] == 168
    document = Document(tmp_path / "reports/epoch-study.docx")
    text = "\n".join(p.text for p in document.paragraphs)
    assert "not supported" in text and "not 168 independent replications" in text
    assert len(document.inline_shapes) == 14
    slides = Presentation(tmp_path / "reports/epoch-study.pptx")
    assert len(slides.slides) == 24
    for slide in slides.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= slides.slide_width
            assert shape.top + shape.height <= slides.slide_height
    for path in (tmp_path / "reports").glob("*.docx"):
        assert zipfile.ZipFile(path).testzip() is None
    assert zipfile.ZipFile(tmp_path / "reports/epoch-study.pptx").testzip() is None
